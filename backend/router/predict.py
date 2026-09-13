from flask import Blueprint, request, jsonify, Response, send_file
from backend.db import get_conn
from backend.services.pipeline import run_pipeline, start_new_session, session_detected_ids
import cv2
import os
import time
import threading
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
import unicodedata

predict_bp = Blueprint("predict", __name__)

# ==============================
# GLOBAL STATE
# ==============================
CURRENT_SESSION = None
VIDEO_SOURCE = None
latest_violations = []

current_stats = {
    "phone": 0,
    "sleep": 0,
    "total": 0,
    "focused": 0,
    "absent": 0,
    "enrolled": 0,
    "visitors": 0, # Số người lạ/khách
    "space_id": None,
    "space_name": "",
    "session_title": ""
}

# Quản lý trạng thái vi phạm thời gian thực (State-based Tracking)
# Cấu trúc: { "student_name_behavior": last_seen_timestamp }
violation_active_states = {}
violation_executor = ThreadPoolExecutor(max_workers=2, thread_name_prefix="violation-io")
JPEG_QUALITY = max(50, min(95, int(os.environ.get("FOCUS_JPEG_QUALITY", "82"))))

# ==============================
# PATH
# ==============================
BASE_DIR = os.path.abspath(os.path.dirname(__file__))
UPLOAD_DIR = os.path.join(BASE_DIR, "..", "web", "static", "uploads")
CAPTURE_DIR = os.path.join(BASE_DIR, "..", "web", "static", "captures")

os.makedirs(UPLOAD_DIR, exist_ok=True)
os.makedirs(CAPTURE_DIR, exist_ok=True)

def fix_frame(frame):
    return frame

# MJPEG STREAM GENERATOR
# ==============================
paused_sessions = set()

@predict_bp.route("/api/session/toggle_pause", methods=["POST"])
def toggle_session_pause():
    data = request.json
    sess_id = data.get("session_id")
    if not sess_id: return jsonify({"error": "No session ID"}), 400
    
    if sess_id in paused_sessions:
        paused_sessions.remove(sess_id)
        status = "resumed"
    else:
        paused_sessions.add(sess_id)
        status = "paused"
    
    return jsonify({"status": "success", "state": status})

class CameraStreamManager:
    def __init__(self):
        self.cap = None
        self.running = False
        self.thread = None
        self.latest_frame_bytes = None
        self.source = None
        self.lock = threading.Lock()
        self.sess_id = None
        self.frame_idx = 0
        self.frame_seq = 0
        self.fps = 30.0
        self.generation = 0

    def start(self, source, sess_id):
        with self.lock:
            if self.running:
                if self.source == source:
                    # Reuse the decoder/AI worker when only the monitoring
                    # session changes. Restarting during inference can leave
                    # multiple workers competing for the same video source.
                    self.sess_id = sess_id
                    print("🔄 Video source already running; session updated in place.")
                    return
                self.stop_unlocked()

            self.source = source
            self.sess_id = sess_id
            self.running = True
            self.generation += 1
            generation = self.generation
            self.frame_idx = 0
            self.latest_frame_bytes = None
            self.thread = threading.Thread(target=self._run, args=(generation,), daemon=True)
            self.thread.start()
            print(f"🎬 Started CameraStreamManager for source: {source}")

    def stop(self):
        with self.lock:
            self.stop_unlocked()

    def stop_unlocked(self):
        self.running = False
        self.generation += 1
        if self.cap:
            try:
                self.cap.release()
            except:
                pass
            self.cap = None
        if self.thread and self.thread.is_alive():
            try:
                self.thread.join(timeout=1.0)
            except:
                pass
        self.thread = None
        self.latest_frame_bytes = None
        print("🛑 Stopped CameraStreamManager thread.")

    def _run(self, generation):
        global CURRENT_SESSION, violation_active_states, current_stats
        
        cap = cv2.VideoCapture(self.source)
        self.cap = cap
        
        if not cap.isOpened():
            print(f"❌ LỖI NGHIÊM TRỌNG: Không thể mở nguồn video '{self.source}'")
            self.running = False
            return

        is_live = False
        try:
            if isinstance(self.source, int) or str(self.source).isdigit() or str(self.source).lower().startswith("rtsp"):
                is_live = True
        except:
            pass

        if is_live:
            # Keep latency low: old camera/RTSP frames are less valuable than the newest one.
            cap.set(cv2.CAP_PROP_BUFFERSIZE, 2)

        fps = cap.get(cv2.CAP_PROP_FPS) or 30.0
        if fps <= 0 or fps > 120:
            fps = 30.0
        self.fps = fps

        while self.running and generation == self.generation and cap.isOpened():
            start_t = time.time()
            if self.sess_id in paused_sessions:
                time.sleep(0.1)
                continue

            success, frame = cap.read()
            if not success:
                cap.set(cv2.CAP_PROP_POS_FRAMES, 0)
                self.frame_idx = 0
                from backend.services.pipeline import reset_session_states
                try:
                    reset_session_states()
                except:
                    pass
                continue

            frame = fix_frame(frame)
            # Publish a raw preview immediately; the processed result replaces
            # it after the first AI pass. This avoids a blank player at startup.
            if self.latest_frame_bytes is None:
                preview_ok, preview_buffer = cv2.imencode(
                    '.jpg', frame, [cv2.IMWRITE_JPEG_QUALITY, JPEG_QUALITY]
                )
                if preview_ok and generation == self.generation:
                    self.latest_frame_bytes = preview_buffer.tobytes()
                    self.frame_seq += 1
            # run_pipeline creates its own enhanced image, so the captured frame
            # can be retained without copying on every iteration.
            clean_frame = frame
            
            self.frame_idx += 1
            virtual_time = None if is_live else (self.frame_idx / fps)
            
            results = []
            from backend.services.pipeline import AI_READY
            if AI_READY:
                results, frame = run_pipeline(frame, session_id=CURRENT_SESSION, current_time=virtual_time)

            if results:
                for res in results:
                    if "Người lạ" in str(res["student_name"]):
                        res["violation"] = True
                        res["behavior"] = "stranger_intrusion"

                p_count = sum(1 for r in results if r["behavior"] == "using_phone")
                s_count = sum(1 for r in results if r["behavior"] == "sleeping")
                f_count = sum(1 for r in results if r["focus"] == "focused")
                v_count = sum(1 for r in results if str(r["student_name"]).startswith("Người lạ"))
                
                current_stats["phone"] = p_count
                current_stats["sleep"] = s_count
                current_stats["total"] = len(results)
                current_stats["focused"] = f_count
                current_stats["visitors"] = v_count
                
                from backend.services.pipeline import session_detected_ids
                total_enrolled = current_stats.get("enrolled", 0)
                current_stats["absent"] = max(0, total_enrolled - len(session_detected_ids))

            if results and self.sess_id:
                now_ts = time.time()
                now_dt = datetime.now()
                
                for res in results:
                    if res["violation"]:
                        s_name = res["student_name"]
                        s_id = res["student_id"]
                        bhv = res["behavior"]
                        
                        if bhv == "stranger_intrusion":
                            incident_key = "stranger_intrusion"
                        else:
                            incident_key = f"{s_name}_{bhv}"
                        
                        is_new_incident = False
                        if incident_key not in violation_active_states:
                            is_new_incident = True
                        else:
                            last_seen = violation_active_states[incident_key]
                            if (now_ts - last_seen > 30):
                                is_new_incident = True
                        
                        violation_active_states[incident_key] = now_ts
                        
                        if is_new_incident:
                            capture_frame = clean_frame.copy()
                            target_is_stranger = "Người lạ" in s_name
                            
                            for r in results:
                                r_is_stranger = "Người lạ" in str(r["student_name"])
                                if r["violation"] and (r_is_stranger == target_is_stranger):
                                    x1, y1, x2, y2 = r["bbox"]
                                    cv2.rectangle(capture_frame, (x1, y1), (x2, y2), (0, 0, 255), 2)
                                    d_n = str(r['student_name']).replace("Người lạ", "NGUOI LA").replace("Đang nhận dạng", "DANG NHAN DANG")
                                    b_l = "PHONE" if r['behavior'] == "using_phone" else ("SLEEP" if r['behavior'] == "sleeping" else "STRANGER")
                                    cv2.putText(capture_frame, f"{d_n} ({b_l})", (x1, y1-10), cv2.FONT_HERSHEY_SIMPLEX, 0.5, (0, 0, 255), 2)

                            safe_name = s_name.replace(' ', '_').replace('#', 'ID')
                            img_name = f"inc_{CURRENT_SESSION}_{int(now_ts)}_{safe_name}.jpg"
                            img_path = os.path.join(CAPTURE_DIR, img_name)
                            img_url = f"/static/captures/{img_name}"

                            def save_violation_worker(session_id, timestamp, b, n, sid, url, path, capture):
                                if session_id is None: return
                                conn = None
                                try:
                                    conn = get_conn()
                                    cursor = conn.cursor()
                                    # Kiểm tra xem phiên đã kết thúc chưa trước khi ghi nhận vi phạm
                                    cursor.execute("SELECT status FROM sessions WHERE id = %s", (session_id,))
                                    sess_row = cursor.fetchone()
                                    if sess_row and sess_row.get("status") == "completed":
                                        cursor.close()
                                        return

                                    cv2.imwrite(path, capture, [cv2.IMWRITE_JPEG_QUALITY, JPEG_QUALITY])
                                    db_sid = None
                                    if sid is not None and str(sid).strip() != "":
                                        try:
                                            db_sid = int(sid)
                                        except:
                                            pass
                                    
                                    if b == 'stranger_intrusion' or "Người lạ" in n:
                                        track_id = 999
                                        try:
                                            if '#' in n: track_id = int(n.split('#')[1])
                                            elif 'ID' in n: track_id = int(n.split('ID')[1])
                                        except: pass
                                        
                                        cursor.execute("""
                                            INSERT INTO stranger_events (session_id, track_id, first_seen, last_seen, image_path, verification_status)
                                            VALUES (%s, %s, %s, %s, %s, 'unverified')
                                        """, (session_id, track_id, timestamp, timestamp, url))
                                    else:
                                        db_behavior = b
                                        if b in ('phone', 'using_phone'): db_behavior = 'using_phone'
                                        elif b in ('sleep', 'sleeping'): db_behavior = 'sleeping'
                                        elif b in ('looking_away', 'turning_away'): db_behavior = 'turning_away'
                                        else: db_behavior = 'unverified'
                                        
                                        cursor.execute("""
                                            INSERT INTO detection_events (session_id, student_id, timestamp, behavior, confidence, image_path)
                                            VALUES (%s, %s, %s, %s, %s, %s)
                                        """, (session_id, db_sid, timestamp, db_behavior, 0.9, url))
                                    conn.commit()

                                    # Notification logic
                                    target_user_id = None
                                    msg = f"⚠️ Phát hiện có người lạ!" if b == 'stranger_intrusion' else f"Học sinh {n} vi phạm: { 'Dùng điện thoại' if b == 'using_phone' else 'Ngủ gật' }"
                                    if db_sid:
                                        cursor.execute("SELECT c.advisor_id FROM users u JOIN classes c ON u.class_id = c.id WHERE u.id = %s AND u.role = 'student'", (db_sid,))
                                        teacher = cursor.fetchone()
                                        if teacher and teacher.get('advisor_id'): target_user_id = teacher['advisor_id']
                                    if not target_user_id:
                                        cursor.execute("SELECT id FROM users WHERE role = 'admin' LIMIT 1")
                                        admin = cursor.fetchone()
                                        target_user_id = admin['id'] if admin else 1
                                    if target_user_id:
                                        cursor.execute("INSERT INTO notifications (user_id, message, created_at) VALUES (%s, %s, %s)", (target_user_id, msg, timestamp))
                                        conn.commit()
                                    cursor.close()
                                except Exception as e_w:
                                    print(f"❌ DB Worker Error: {e_w}")
                                finally:
                                    if conn: conn.close()

                            violation_executor.submit(
                                save_violation_worker,
                                self.sess_id, now_dt, bhv, s_name, s_id,
                                img_url, img_path, capture_frame,
                            )

                            latest_violations.insert(0, {
                                "id": str(int(now_ts * 1000)) + "_" + str(res['id']),
                                "url": img_url,
                                "time": now_dt.strftime("%H:%M:%S"),
                                "behavior": bhv,
                                "student_name": s_name,
                                "student_code": res.get("student_code", ""),
                                "space_name": current_stats["space_name"],
                                "individual_focus_score": res.get("individual_focus_score", 0),
                                "violation_duration": res.get("violation_duration", 0)
                            })
                            if len(latest_violations) > 50: latest_violations.pop()

            encoded, buffer = cv2.imencode(
                '.jpg', frame, [cv2.IMWRITE_JPEG_QUALITY, JPEG_QUALITY]
            )
            if encoded and generation == self.generation:
                self.latest_frame_bytes = buffer.tobytes()
                self.frame_seq += 1
            
            # Đồng bộ thời gian thực: Trừ đi thời gian xử lý AI để đảm bảo luồng video mượt như offline
            elapsed = time.time() - start_t
            sleep_time = max(0.001, (1.0 / self.fps) - elapsed)
            time.sleep(sleep_time)

        cap.release()
        if generation == self.generation:
            self.cap = None
            self.running = False

stream_manager = CameraStreamManager()

def generate_frames(sess_id, video_source):
    global stream_manager
    stream_manager.start(video_source, sess_id)
    last_seq = -1
    while stream_manager.running:
        frame_bytes = stream_manager.latest_frame_bytes
        frame_seq = stream_manager.frame_seq
        if frame_bytes is not None and frame_seq != last_seq:
            last_seq = frame_seq
            yield (b'--frame\r\nContent-Type: image/jpeg\r\n\r\n' + frame_bytes + b'\r\n')
        else:
            time.sleep(0.01)

@predict_bp.route("/stream_video/<room_id>")
def stream_video(room_id):
    global VIDEO_SOURCE, current_stats
    conn = None
    try:
        conn = get_conn()
        cursor = conn.cursor()
        
        space = None
        room_str = str(room_id).strip()
        if room_str.isdigit():
            cursor.execute("SELECT id, name, video_path FROM spaces WHERE id = %s", (int(room_str),))
            space = cursor.fetchone()
        if not space:
            cursor.execute("SELECT id, name, video_path FROM spaces WHERE name = %s", (room_str,))
            space = cursor.fetchone()
        if not space:
            cursor.execute("SELECT id, name, video_path FROM spaces ORDER BY id DESC LIMIT 1")
            space = cursor.fetchone()

        if not space:
            return jsonify({"error": "Space not found"}), 404
        
        space_id = space["id"]
        current_stats["space_name"] = space["name"]
        v_path = str(space["video_path"])
        
        if os.path.exists(v_path):
            source = v_path
        else:
            upload_candidate = os.path.join(UPLOAD_DIR, os.path.basename(v_path))
            if os.path.exists(upload_candidate):
                source = upload_candidate
            else:
                source = VIDEO_SOURCE or 0
        
        cursor.execute("""
            SELECT s.id FROM sessions s
            JOIN course_sections cs ON s.section_id = cs.id
            WHERE cs.space_id = %s AND s.status = 'ongoing' 
            ORDER BY s.id DESC LIMIT 1
        """, (space_id,))
        sess = cursor.fetchone()
        sess_id = sess["id"] if sess else None
        
        if not sess_id:
            print(f"⚠️ Warning: No active session for room {space_id}. Violations will NOT be saved.")
            
        cursor.close()
        return Response(generate_frames(sess_id, source), mimetype='multipart/x-mixed-replace; boundary=frame')
    except Exception as e:
        print(f"❌ STREAM VIDEO ERROR: {e}")
        return jsonify({"error": str(e)}), 500
    finally:
        if conn:
            conn.close()

@predict_bp.route("/get_realtime_data", methods=["GET"])
def get_realtime_data():
    return jsonify({"counts": current_stats, "history": latest_violations})

@predict_bp.route("/select_room", methods=["POST", "GET"])
def select_room():
    global VIDEO_SOURCE, CURRENT_SESSION, latest_violations, current_stats, violation_active_states
    conn = None
    try:
        latest_violations = []
        violation_active_states = {}
        current_stats = {"phone": 0, "sleep": 0, "total": 0, "focused": 0, "absent": 0, "enrolled": 0, "visitors": 0, "space_id": None, "space_name": ""}
        
        data = request.get_json(silent=True) or request.form or request.args or {}
        room_id = data.get("room") or data.get("room_id") or data.get("space_id")
        
        conn = get_conn()
        cursor = conn.cursor()
        
        space = None
        if room_id is not None:
            room_str = str(room_id).strip()
            if room_str.isdigit():
                cursor.execute("SELECT id, name, video_path FROM spaces WHERE id = %s", (int(room_str),))
                space = cursor.fetchone()
            if not space:
                cursor.execute("SELECT id, name, video_path FROM spaces WHERE name = %s", (room_str,))
                space = cursor.fetchone()

        if not space:
            # Fallback về phòng học mới nhất trong DB
            cursor.execute("SELECT id, name, video_path FROM spaces ORDER BY id DESC LIMIT 1")
            space = cursor.fetchone()

        if not space:
            return jsonify({"error": "No spaces available in system"}), 400

        v_path = str(space["video_path"])
        room_name = space["name"]
        space_id = space["id"]
        
        if os.path.exists(v_path):
            VIDEO_SOURCE = v_path
        else:
            VIDEO_SOURCE = os.path.join(UPLOAD_DIR, os.path.basename(v_path))

        cursor.execute("""
            SELECT s.id, s.title 
            FROM sessions s
            JOIN course_sections cs ON s.section_id = cs.id
            WHERE cs.space_id = %s AND s.status = 'ongoing' 
            AND s.start_time > DATE_SUB(NOW(), INTERVAL 2 HOUR)
            ORDER BY s.id DESC LIMIT 1
        """, (space_id,))
        existing_sess = cursor.fetchone()
        
        if existing_sess:
            CURRENT_SESSION = existing_sess["id"]
            title = existing_sess["title"]
            print(f"🔄 Reusing existing session: {CURRENT_SESSION}")
        else:
            cursor.execute("""
                SELECT c.title AS subject_title 
                FROM course_sections cs
                JOIN courses c ON cs.course_id = c.id
                WHERE cs.space_id = %s LIMIT 1
            """, (space_id,))
            sch = cursor.fetchone()
            title = sch["subject_title"] if sch else f"Môn học {room_name}"
            CURRENT_SESSION = start_new_session(title=title, space_id=space_id)
            print(f"🆕 Started NEW session: {CURRENT_SESSION}")

        current_stats["space_id"] = space_id
        current_stats["space_name"] = room_name
        
        from backend.services.pipeline import get_current_stats, session_detected_ids
        stats = get_current_stats(CURRENT_SESSION)
        final_enrolled = stats["total_enrolled"]
        
        current_stats["enrolled"] = final_enrolled
        current_stats["absent"] = max(0, final_enrolled - len(session_detected_ids))
        current_stats["session_title"] = title
        return jsonify({"status": "success", "session_id": CURRENT_SESSION, "session_title": title, "start_time": datetime.now().strftime("%H:%M:%S"), "session_status": "ongoing", "counts": current_stats})
    except Exception as e:
        print(f"❌ SELECT ROOM ERROR: {e}")
        return jsonify({"error": str(e)}), 400
    finally:
        if conn:
            conn.close()

@predict_bp.route("/api/session/update_siso", methods=["POST"])
def update_session_siso():
    global current_stats, CURRENT_SESSION
    data = request.json
    new_siso = data.get("siso")
    if not CURRENT_SESSION or new_siso is None:
        return jsonify({"error": "No active session or missing data"}), 400
    
    # Cập nhật sĩ số thủ công trực tiếp trong RAM
    current_stats["enrolled"] = int(new_siso)
    from backend.services.pipeline import session_detected_ids
    current_stats["absent"] = max(0, int(new_siso) - len(session_detected_ids))
    
    return jsonify({"status": "success", "new_siso": new_siso})

@predict_bp.route("/api/sessions/<int:session_id>/export_attendance", methods=["GET"])
def export_attendance(session_id):
    from backend.services.pipeline import session_detected_ids
    conn = None
    try:
        conn = get_conn()
        cursor = conn.cursor()
        
        cursor.execute("""
            SELECT cs.space_id, s.title, s.section_id 
            FROM sessions s
            JOIN course_sections cs ON s.section_id = cs.id
            WHERE s.id = %s
        """, (session_id,))
        session_info = cursor.fetchone()
        if not session_info: return jsonify({"error": "Session not found"}), 404
        section_id = session_info['section_id']
        
        cursor.execute("""
            SELECT u.id, u.name, u.code AS student_code 
            FROM users u 
            JOIN section_enrollments se ON u.id = se.student_id 
            WHERE se.section_id = %s AND u.role = 'student'
        """, (section_id,))
        students = cursor.fetchall()
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Diem Danh"
        
        # Merge cells & titles
        ws.merge_cells('A1:G1')
        ws['A1'] = "DANH SÁCH ĐIỂM DANH NHÓM 01_HK2_25-26"
        ws['A1'].font = Font(bold=True, size=16)
        ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
        
        title_str = session_info['title'] or ""
        norm_title = unicodedata.normalize('NFC', title_str.lower())
        code_map = {
            "nosql": "CSDL_NOSQL",
            "mobile": "MOB104",
            "trí tuệ": "AI101",
            "python": "PY201",
            "xử lý ảnh": "DIP301",
            "blockchain": "BC401",
            "an ninh mạng": "SEC501",
            "bảo mật": "SEC501",
            "fullstack": "WEB202",
            "web": "WEB202",
            "tự học": "TU_HOC",
            "thư viện": "THU_VIEN"
        }
        derived_code = "CHUYEN_DE"
        for k, v in code_map.items():
            if unicodedata.normalize('NFC', k) in norm_title:
                derived_code = v
                break
                
        ws['A3'] = f"Môn: {title_str}"
        ws['B3'] = f"Mã môn: {derived_code}"
        ws['C3'] = "Năm Học: 25-26"
        
        ws['E4'] = "Ngày Học:"
        ws['E5'] = "Buổi"
        
        ws['F4'] = datetime.now().strftime("%d/%m/%Y")
        
        headers = ["Stt", "Mã SV", "Họ SV", "Ngày Sinh", "Lớp", "01", "02"]
        thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
        fill = PatternFill(start_color="E2E2E2", end_color="E2E2E2", fill_type="solid")
        
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=6, column=col_num, value=header)
            cell.font = Font(bold=True)
            cell.fill = fill
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = thin_border
            
        REAL_NAMES = {
            "22050001": "Đặng Ngọc Phong", "22050076": "Hạ Văn Minh",
            "22050025": "Huỳnh Minh Chiến", "22050086": "Lê Đình Quốc",
            "22050089": "Phạm Hồng Quý", "22050044": "Phan Văn Lộc",
            "22050034": "Quách Thị Thu", "22050015": "Thiều Đăng Hùng",
            "22050023": "Tô Gia Dân", "22050035": "Trần Văn Tài",
            "22050038": "Vũ Duy Hoàng", "22050059": "Vương Quốc Cường"
        }
        
        for i, s in enumerate(students, 1):
            row_idx = 6 + i
            real_name = REAL_NAMES.get(str(s['student_code']), s['name'])
            
            ws.cell(row=row_idx, column=1, value=i).alignment = Alignment(horizontal='center')
            ws.cell(row=row_idx, column=2, value=s['student_code']).alignment = Alignment(horizontal='center')
            ws.cell(row=row_idx, column=3, value=real_name)
            ws.cell(row=row_idx, column=4, value="11/04/2004").alignment = Alignment(horizontal='center')
            ws.cell(row=row_idx, column=5, value="25TH01").alignment = Alignment(horizontal='center')
            
            if s['id'] in session_detected_ids:
                ws.cell(row=row_idx, column=6, value="Có").alignment = Alignment(horizontal='center')
                
            for col in range(1, 8):
                ws.cell(row=row_idx, column=col).border = thin_border
                
        # Set column widths
        ws.column_dimensions['B'].width = 15
        ws.column_dimensions['C'].width = 25
        ws.column_dimensions['D'].width = 15
        ws.column_dimensions['E'].width = 15
        
        file_path = os.path.join(UPLOAD_DIR, f"DiemDanh_{session_id}.xlsx")
        wb.save(file_path)
        
        return send_file(file_path, as_attachment=True, download_name=f"DanhSachDiemDanh_{datetime.now().strftime('%d%m%Y')}.xlsx")
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    finally:
        if conn: conn.close()

@predict_bp.route("/api/sessions/<int:session_id>/attendance_data", methods=["GET"])
def get_attendance_data(session_id):
    from backend.services.pipeline import session_detected_ids
    conn = None
    try:
        conn = get_conn()
        cursor = conn.cursor()
        cursor.execute("SELECT section_id FROM sessions WHERE id = %s", (session_id,))
        sess = cursor.fetchone()
        if not sess: return jsonify({"error": "Session not found"}), 404
        section_id = sess['section_id']
        
        cursor.execute("""
            SELECT u.id, u.name, u.code AS student_code 
            FROM users u JOIN section_enrollments se ON u.id = se.student_id 
            WHERE se.section_id = %s AND u.role = 'student'
        """, (section_id,))
        students = cursor.fetchall()
        
        REAL_NAMES = {
            "22050001": "Đặng Ngọc Phong", "22050076": "Hạ Văn Minh",
            "22050025": "Huỳnh Minh Chiến", "22050086": "Lê Đình Quốc",
            "22050089": "Phạm Hồng Quý", "22050044": "Phan Văn Lộc",
            "22050034": "Quách Thị Thu", "22050015": "Thiều Đăng Hùng",
            "22050023": "Tô Gia Dân", "22050035": "Trần Văn Tài",
            "22050038": "Vũ Duy Hoàng", "22050059": "Vương Quốc Cường"
        }
        
        res = []
        for i, s in enumerate(students, 1):
            res.append({
                "stt": i,
                "code": s['student_code'],
                "name": REAL_NAMES.get(str(s['student_code']), s['name']),
                "dob": "11/04/2004",
                "class_name": "25TH01",
                "present": s['id'] in session_detected_ids
            })
        return jsonify(res)
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    finally:
        if conn: conn.close()
